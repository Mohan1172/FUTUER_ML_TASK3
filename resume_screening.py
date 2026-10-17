"""Core resume parsing, skill matching, and ranking logic.

This module deliberately avoids demographic, protected, or other sensitive attributes.
It provides transparent decision-support scores, not a hiring decision.
"""
from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Iterable, Sequence

import numpy as np
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

DEFAULT_SKILLS = [
    "python", "java", "c++", "c#", "c", "sql", "mysql", "postgresql",
    "scikit-learn", "machine learning", "deep learning", "natural language processing",
    "nlp", "pandas", "numpy", "tensorflow", "pytorch", "data analysis",
    "data visualization", "excel", "power bi", "tableau", "statistics",
    "git", "github", "docker", "kubernetes", "aws", "azure", "gcp",
    "rest api", "api development", "flask", "django", "fastapi", "streamlit",
    "html", "css", "javascript", "typescript", "react", "node.js",
    "linux", "networking", "cybersecurity", "testing", "selenium",
    "agile", "communication", "problem solving", "project management",
    "computer vision", "data structures", "algorithms", "mongodb",
    "data engineering", "etl", "spark", "hadoop", "jupyter notebook",
]

def normalize_text(text: str) -> str:
    """Normalize whitespace while preserving useful punctuation in technical skills."""
    text = (text or "").replace("\x00", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()

def normalize_skill_list(raw: str | Iterable[str]) -> list[str]:
    """Turn a comma/newline separated skill list into unique lowercase skill names."""
    if isinstance(raw, str):
        items = re.split(r"[,;\n]", raw)
    else:
        items = list(raw)
    seen = set()
    cleaned = []
    for item in items:
        skill = normalize_text(str(item)).strip(" .\t").lower()
        if skill and skill not in seen:
            seen.add(skill)
            cleaned.append(skill)
    return cleaned

def extract_text_from_upload(filename: str, content: bytes) -> str:
    """Extract text from a PDF, DOCX, or TXT byte stream."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".txt":
        for encoding in ("utf-8", "utf-8-sig", "latin-1"):
            try:
                return normalize_text(content.decode(encoding))
            except UnicodeDecodeError:
                continue
        return normalize_text(content.decode("utf-8", errors="replace"))
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        return normalize_text(" ".join(page.extract_text() or "" for page in reader.pages))
    if suffix == ".docx":
        from docx import Document
        document = Document(io.BytesIO(content))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" ".join(cell.text for cell in row.cells))
        return normalize_text(" ".join(parts))
    raise ValueError("Unsupported file type. Use PDF, DOCX, or TXT.")

def _skill_present(skill: str, text: str) -> bool:
    """Match whole skill phrases to reduce false positives (e.g. 'C' inside 'computer')."""
    skill = skill.lower().strip()
    text = text.lower()
    # Allow flexible whitespace for multi-word skill phrases.
    escaped = re.escape(skill)
    escaped = escaped.replace(r"\ ", r"\s+")
    # Word boundaries work for most terms; technical tokens such as C++ are supported.
    pattern = rf"(?<![a-z0-9]){escaped}(?![a-z0-9])"
    return re.search(pattern, text, flags=re.IGNORECASE) is not None

def extract_skills(text: str, skills: Sequence[str]) -> list[str]:
    normalized = normalize_text(text)
    return [skill for skill in skills if _skill_present(skill, normalized)]

def screen_resumes(
    job_description: str,
    resumes: Sequence[dict],
    skills: Sequence[str] | None = None,
    required_skills: Sequence[str] | None = None,
    preferred_skills: Sequence[str] | None = None,
    skill_weight: float = 0.65,
):
    """Score and rank resumes using dictionary skill coverage and TF-IDF cosine similarity.

    overall = skill_weight * skill_match + (1 - skill_weight) * text_similarity.
    Skill match is coverage of the skill dictionary skills found in the job description.
    If no dictionary skills occur in the job description, it falls back to all supplied skills.
    """
    import pandas as pd

    job_description = normalize_text(job_description)
    if not job_description:
        raise ValueError("Job description cannot be empty.")
    if not resumes:
        raise ValueError("At least one resume is required.")

    skills = normalize_skill_list(skills or DEFAULT_SKILLS)
    required_skills = normalize_skill_list(required_skills or [])
    preferred_skills = normalize_skill_list(preferred_skills or [])
    skill_weight = min(1.0, max(0.0, float(skill_weight)))

    job_skills = extract_skills(job_description, skills)
    target_skills = job_skills if job_skills else skills
    required_in_scope = required_skills
    preferred_in_scope = preferred_skills

    documents = [job_description] + [normalize_text(item.get("text", "")) for item in resumes]
    vectorizer = TfidfVectorizer(stop_words="english", ngram_range=(1, 2), sublinear_tf=True)
    try:
        matrix = vectorizer.fit_transform(documents)
        similarities = cosine_similarity(matrix[1:], matrix[0]).ravel()
    except ValueError:
        # Happens when all text is empty or contains only stop words.
        similarities = np.zeros(len(resumes), dtype=float)

    results = []
    for index, item in enumerate(resumes):
        resume_text = normalize_text(item.get("text", ""))
        found = extract_skills(resume_text, skills)
        target_found = [s for s in target_skills if s in found]
        target_missing = [s for s in target_skills if s not in found]
        skill_match = len(target_found) / len(target_skills) if target_skills else 0.0

        required_found = [s for s in required_in_scope if _skill_present(s, resume_text)]
        required_missing = [s for s in required_in_scope if s not in required_found]
        preferred_found = [s for s in preferred_in_scope if _skill_present(s, resume_text)]
        preferred_missing = [s for s in preferred_in_scope if s not in preferred_found]

        similarity = float(similarities[index]) if index < len(similarities) else 0.0
        score = 100.0 * (skill_weight * skill_match + (1.0 - skill_weight) * similarity)
        results.append({
            "filename": item.get("filename", f"resume_{index + 1}"),
            "overall_score": round(score, 2),
            "skill_match_pct": round(skill_match * 100.0, 2),
            "text_similarity_pct": round(similarity * 100.0, 2),
            "job_skills_detected": job_skills,
            "matched_skills": target_found,
            "missing_skills": target_missing,
            "required_skills_found": required_found,
            "required_skills_missing": required_missing,
            "preferred_skills_found": preferred_found,
            "preferred_skills_missing": preferred_missing,
        })

    results.sort(key=lambda row: (-row["overall_score"], row["filename"].lower()))
    for rank, row in enumerate(results, start=1):
        row["rank"] = rank
    columns = [
        "rank", "filename", "overall_score", "skill_match_pct", "text_similarity_pct",
        "job_skills_detected", "matched_skills", "missing_skills",
        "required_skills_found", "required_skills_missing",
        "preferred_skills_found", "preferred_skills_missing",
    ]
    return pd.DataFrame(results, columns=columns)
